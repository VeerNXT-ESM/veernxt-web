"""
Build a self-contained, profiler-ready workbook from the generated personas:
docs/Profiler_Test_Input_<date>.xlsx  (untracked, like other docs/*.xlsx)

Every profile column is named exactly like a field of `profileSchema`
(api/profile/recommend.js), so a row can be POSTed to /api/profile/recommend as-is by
scripts/testdata/run_profiler_from_xlsx.py. Personas are anonymised (synthetic name, email,
phone, DOB); no real respondent data is in the file.

Sheets:
  README                       what this is and how to run it
  Profiles - full survey data  everything the survey knows, mapped to the profiler's vocabulary
  Profiles - live form only    only what the live form can send (the rest = defaults)
  Field guide                  every field: allowed values, whether the engine uses it, how it was filled

Usage:  python scripts/testdata/build_personas.py && python scripts/testdata/make_profiler_input_xlsx.py
"""
import datetime as dt
import json
import pathlib

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = pathlib.Path(__file__).resolve().parents[2]
PERSONAS = ROOT / "test-data" / "personas"
OUT = ROOT / "docs" / f"Profiler_Test_Input_{dt.date.today().isoformat()}.xlsx"

# Field order and types mirror profileSchema in api/profile/recommend.js
FIELDS = [
    "fullName", "dateOfBirth", "category", "disabilityStatus", "disabilityType", "disabilityPercentage",
    "stateOfDomicile", "district", "maritalStatus", "email", "mobile",
    "serviceBranch", "armCorpsTrade", "roleAppointment", "totalServiceDuration", "militaryCourses",
    "characterOnDischarge", "specificSkills",
    "highestQualification", "completedDuringService", "nccCertification", "sportsAchievement", "mathInClass12",
    "heightCm", "weightKg", "chestCm", "chestExpansion", "vision", "colourBlind", "medicalCategory", "physicalProficiency",
    "careerPreferences", "relocation", "englishComfort", "sewaNidhiInterests", "consent",
]
ARRAYS = {"militaryCourses", "specificSkills", "careerPreferences", "sewaNidhiInterests"}
BOOLS = {"completedDuringService", "mathInClass12", "colourBlind", "consent"}
TEXT = {"dateOfBirth", "mobile"}  # keep as text: no Excel date conversion, keep leading zeros
SEP = " | "

ENUMS = {
    "category": ["General", "OBC", "SC", "ST", "EWS"],
    "disabilityStatus": ["No", "Yes"],
    "disabilityType": ["Locomotor", "Visual", "Hearing", "Intellectual", "Multiple", "Other"],
    "disabilityPercentage": ["40-49%", "50-69%", "70-100%"],
    "maritalStatus": ["Single", "Married"],
    "serviceBranch": ["Indian Army", "Indian Navy", "Indian Air Force"],
    "characterOnDischarge": ["Exemplary", "Very Good", "Good"],
    "highestQualification": ["Class 10", "Class 12", "Graduate", "Post-Graduate"],
    "nccCertification": ["None", "A Certificate", "B Certificate", "C Certificate"],
    "sportsAchievement": ["None", "District", "State", "National", "International/Services"],
    "physicalProficiency": ["Excellent", "Good", "Satisfactory"],
    "relocation": ["Home District", "Home State", "Anywhere in India"],
    "englishComfort": ["Basic", "Intermediate", "Fluent"],
}
REQUIRED_ENUMS = {"category", "disabilityStatus", "maritalStatus", "serviceBranch", "characterOnDischarge", "highestQualification"}

# (field, required by API, allowed values / format, used by the engine?, asked by live form?, in survey?, how filled)
GUIDE = [
    ("fullName", "yes", "text, min 2 chars", "no (echoed in summary only)", "yes", "yes", "Replaced by 'Persona <id>' (anonymised)"),
    ("dateOfBirth", "yes", "YYYY-MM-DD", "NO - never used; there is no age check", "yes (age 18-50)", "yes", "Regenerated: same age as the respondent, random day. Age is not used by the engine."),
    ("category", "yes", "General / OBC / SC / ST / EWS", "yes: small reservation bonus (SC/ST 5, OBC/EWS 3)", "yes", "yes", "Survey 'OBC (Non-creamy layer)' -> OBC; 'General (EWS)' -> EWS"),
    ("disabilityStatus", "yes", "No / Yes", "NO - stored only", "yes", "no", "Assumed 'No' for survey people; edge personas X12/X13 set it"),
    ("disabilityType", "no", "Locomotor / Visual / Hearing / Intellectual / Multiple / Other", "NO - stored only", "yes", "no", "Blank unless disabilityStatus = Yes"),
    ("disabilityPercentage", "no", "40-49% / 50-69% / 70-100%", "NO - stored only", "yes", "no", "Blank unless disabilityStatus = Yes"),
    ("stateOfDomicile", "yes", "State/UT name as in src/lib/districts.js", "YES: state exams and home-state bonus", "yes", "yes", "Normalised to the app's spellings (e.g. 'Jammu & Kashmir')"),
    ("district", "no", "text", "no", "yes", "yes", "Title-cased survey value"),
    ("maritalStatus", "yes", "Single / Married", "no", "yes", "yes", "As given (all 'Single' in survey)"),
    ("email", "yes", "valid email", "no", "yes", "yes", "Synthetic: veernxt-test+<id>@example.com"),
    ("mobile", "yes", "7-15 digits", "no", "yes", "yes", "Synthetic: 0000<6 digits>"),
    ("serviceBranch", "yes", "Indian Army / Indian Navy / Indian Air Force", "no (summary only)", "yes", "yes", "As given (all Army in survey)"),
    ("armCorpsTrade", "yes", "arm/corps name from designations.json, or a trade word", "YES: trade -> career tracks", "yes (dropdown)", "yes (free text)", "Survey 'Armd/Armoured/Armd gd' -> ARMOURED CORPS; 'EME' -> EME; unresolved -> 'Unspecified'"),
    ("roleAppointment", "yes", "role text", "YES: trade -> career tracks", "yes (dropdown from designations.json)", "yes (free text)", "Abbreviations expanded (GD -> General Duty, Dvr -> Driver); NOT in the live dropdown for Armoured Corps"),
    ("totalServiceDuration", "yes", "'<n> years <m> months'", "yes: 4+ years gives full-term bonus", "yes", "NO", "ASSUMED 4 years 0 months for survey people (Agniveer)"),
    ("militaryCourses", "no", "list, separated by ' | '", "no", "NO", "yes", "Junk answers (No/Nil/None) removed"),
    ("characterOnDischarge", "yes", "Exemplary / Very Good / Good", "yes: bonus, larger for Banking/PSU/Police/Judiciary", "yes", "yes", "As given"),
    ("specificSkills", "no", "list, separated by ' | ' (engine keys, e.g. 'Weapons Handling', 'Driving (LMV/HMV)')", "yes: skills -> career tracks", "NO", "yes", "Survey wording mapped to the exact keys the engine looks up"),
    ("highestQualification", "yes", "Class 10 / Class 12 / Graduate / Post-Graduate", "YES: hard qualification gate + fit bonus", "yes", "yes", "Matriculation -> Class 10; Intermediate -> Class 12; Graduation -> Graduate"),
    ("completedDuringService", "no", "TRUE / FALSE", "no", "NO (default FALSE)", "yes (in the qualification answer)", "TRUE only for 'Graduation while in service'"),
    ("nccCertification", "no", "None / A / B / C Certificate", "yes: bonus for exams flagged ncc_bonus", "yes", "yes", "'N/A' -> None"),
    ("sportsAchievement", "no", "None / District / State / National / International/Services", "yes: bonus for police/defence/railways/sports-quota exams", "NO (default None)", "yes", "'Services' -> International/Services"),
    ("mathInClass12", "no", "TRUE / FALSE", "yes: +8 / -10 on maths-required exams", "NO (default FALSE)", "NO", "ASSUMED FALSE for survey people"),
    ("heightCm", "yes", "number 100-250", "no (stored only)", "yes", "yes", "Cleaned ('170cm' -> 170); unusable values imputed with the cohort median"),
    ("weightKg", "no", "number 30-200", "no (stored only)", "yes", "NO", "Blank for survey people (not collected)"),
    ("chestCm", "no", "number", "no (stored only)", "yes", "yes", "First number of '82 & 88'; blank if unusable"),
    ("chestExpansion", "no", "number", "no (stored only)", "yes", "yes (combined with chest)", "Second number minus first, or the stated expansion; blank if unclear"),
    ("vision", "no", "text", "no", "NO", "yes", "As given"),
    ("colourBlind", "no", "TRUE / FALSE", "NO - stored only", "NO (default FALSE)", "no", "FALSE (edge persona X13 sets TRUE)"),
    ("medicalCategory", "no", "SHAPE-1 / SHAPE-2 / ...", "YES: anything but SHAPE-1 removes all physical-required exams", "NO (default SHAPE-1)", "yes", "'Good' assumed SHAPE-1 (ambiguous); 'F(2)p' -> SHAPE-2"),
    ("physicalProficiency", "no", "Excellent / Good / Satisfactory", "yes: 'Satisfactory' drops police/defence physical exams", "NO (default Good)", "yes", "As given"),
    ("careerPreferences", "yes (min 1)", "list, separated by ' | '", "YES: bonus by preference", "yes (form uses track codes)", "yes (Central/State/Private/Entrepreneurship)", "'full survey' sheet: Central Government / State Government / Private Sector / Entrepreneurship. 'live form' sheet: the form's track codes via an assumed mapping. NOTE: the scorer only recognises the words central/state/bank/psu/private/entrepren"),
    ("relocation", "no", "Home District / Home State / Anywhere in India", "yes: 'Anywhere in India' opens other states' exams (not used for Central exams)", "NO (default Home State)", "yes", "'Pan India' -> Anywhere in India"),
    ("englishComfort", "no", "Basic / Intermediate / Fluent", "yes: bonus; penalty for English-heavy exams if Basic", "NO (default Basic)", "yes", "As given"),
    ("sewaNidhiInterests", "no", "list, separated by ' | '", "no", "yes (different options)", "yes (different options)", "Survey value kept in the full sheet; blank in the live-form sheet"),
    ("consent", "yes (must be TRUE)", "TRUE", "no", "yes", "yes", "TRUE"),
]

HEAD = PatternFill("solid", fgColor="1F3A2E")
NOTE = PatternFill("solid", fgColor="E8EFE9")
EXTRA = PatternFill("solid", fgColor="6B7280")


def cell_value(field, v):
    if field in ARRAYS:
        return SEP.join(v) if v else None
    if v in (None, ""):
        return None
    return v


def profiles_sheet(wb, title, variant, personas):
    ws = wb.create_sheet(title)
    header = FIELDS + ["persona_id", "source", "purpose", "cleaned_or_assumed"]
    ws.append(header)
    for i, h in enumerate(header, 1):
        c = ws.cell(1, i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD if h in FIELDS else EXTRA
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for p in personas:
        prof = p[variant]
        row = [cell_value(f, prof.get(f)) for f in FIELDS]
        row += [p["id"], p["source"], p.get("purpose", ""), ", ".join(p.get("flags", []))]
        ws.append(row)
    n = ws.max_row
    for ci, f in enumerate(FIELDS, 1):
        col = get_column_letter(ci)
        if f in TEXT:
            for r in range(2, n + 1):
                ws.cell(r, ci).number_format = "@"
        if f in ENUMS:
            dv = DataValidation(type="list", formula1='"' + ",".join(ENUMS[f]) + '"', allow_blank=f not in REQUIRED_ENUMS)
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}{n}")
        if f in BOOLS:
            dv = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}{n}")
    widths = {"fullName": 16, "dateOfBirth": 12, "stateOfDomicile": 20, "armCorpsTrade": 30, "roleAppointment": 20, "email": 34,
              "careerPreferences": 40, "specificSkills": 40, "militaryCourses": 22, "purpose": 50, "cleaned_or_assumed": 60,
              "totalServiceDuration": 20, "sewaNidhiInterests": 22}
    for ci, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(ci)].width = widths.get(h, 14)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def main():
    personas = json.loads((PERSONAS / "survey_personas.json").read_text(encoding="utf-8")) + \
        json.loads((PERSONAS / "synthetic_personas.json").read_text(encoding="utf-8"))
    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    lines = [
        ("Profiler test input", True),
        (f"Generated {dt.date.today().isoformat()} from the survey export plus hand-built edge cases. {len(personas)} personas.", False),
        ("", False),
        ("What this is", True),
        ("Each row is a complete profile whose column names are exactly the fields the profiler validates (profileSchema in api/profile/recommend.js), so it can be sent to /api/profile/recommend without changes.", False),
        ("Two sheets hold the same people: 'Profiles - full survey data' (everything the survey knows) and 'Profiles - live form only' (only what the live form can send; other fields are defaults). Comparing the two runs shows what the form's missing questions cost.", False),
        ("Grey columns at the far right (persona_id, source, purpose, cleaned_or_assumed) are for humans; the runner ignores them.", False),
        ("Anonymised: names, emails, phones and DOBs are synthetic. No real respondent data is in this file.", False),
        ("", False),
        ("How to run (read-only)", True),
        ("python scripts/testdata/run_profiler_from_xlsx.py --file <this file> --url https://www.veernxt.in/api/profile/recommend", False),
        ("(use the www host: veernxt.in answers with a 307 redirect, which a POST does not follow. For a local copy use http://localhost:3000/api/profile/recommend.)", False),
        ("Needs Python 3 and openpyxl (pip install openpyxl). The runner sends NO login token, so the profiler never writes a user record or awards points; its only database access is reading the exams table.", False),
        ("It writes a results CSV: for each person the number of eligible exams and the top 10 recommendations (exam, level, state, score).", False),
        ("Add --sheet 'Profiles - live form only' to run the other variant, or --limit 5 for a quick smoke test.", False),
        ("", False),
        ("Editing", True),
        ("Enum columns have dropdowns. Lists use ' | ' as the separator. Dates are YYYY-MM-DD text. Keep mobile as text. See the 'Field guide' sheet for every field, whether the engine uses it, and how it was filled.", False),
        ("", False),
        ("Fields the engine does NOT use (they pass validation but do not change results)", True),
        ("dateOfBirth (no age check), disability fields, colourBlind, height/weight/chest, vision, marital status, district, serviceBranch, courses, sewaNidhiInterests.", False),
    ]
    for text, bold in lines:
        readme.append([text])
        readme.cell(readme.max_row, 1).font = Font(bold=bold, size=13 if bold else 11)
        readme.cell(readme.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    readme.column_dimensions["A"].width = 140

    profiles_sheet(wb, "Profiles - full survey data", "profile_survey", personas)
    profiles_sheet(wb, "Profiles - live form only", "profile_app", personas)

    g = wb.create_sheet("Field guide")
    g.append(["Field", "Required by API", "Allowed values / format", "Used by the engine?", "Asked by the live form?", "In the survey?", "How it was filled"])
    for c in g[1]:
        c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), HEAD, Alignment(wrap_text=True, vertical="center")
    for row in GUIDE:
        g.append(list(row))
    for i, w in enumerate([22, 16, 46, 44, 26, 26, 80], 1):
        g.column_dimensions[get_column_letter(i)].width = w
    for row in g.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if str(row[3].value).startswith(("NO", "no")):
            row[3].fill = NOTE
    g.freeze_panes = "B2"
    wb.save(OUT)
    print(f"wrote {OUT} ({len(personas)} personas x 2 sheets)")


if __name__ == "__main__":
    main()
